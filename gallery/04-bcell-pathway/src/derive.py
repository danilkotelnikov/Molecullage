"""Derive the numbers of the figure from the paper's supplementary tables (data/, see fetch_data.py).

    py -3 -B -I gallery/04-bcell-pathway/src/derive.py

Reads Table S1 (47 reporter genes), S3A/S3B (activators, repressors), S5 (genes changed by rapamycin), S6A/S6B (STRING clusters),
S9 (genes changed by the GSK3 inhibitor LY2090314) and S10 (genes changed by both). Writes small derived tables to tables/:
    reporters_47.csv   the 47 reporter genes, their group in the figure, the name in Table S1 and Table S3, activators and repressors per screen
    counts.json        every count the figure prints, with the table or text it comes from
    example_genes.csv  log2 fold change of the genes the paper names, in S5, S9 and S10 (empty: not listed in that table)
    shared_303.csv     the 303 genes listed in both S5 and S9: log2 fold change under rapamycin and under LY2090314
Every count that the paper's text also states is asserted against the text. A mismatch stops the script.
"""
from __future__ import annotations

import csv
import json
import math
from collections import Counter
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
DATA, OUT = ROOT / "data", ROOT / "tables"

# Groups of the 47 reporter genes as drawn in the paper's Fig. 2A (the text names the four groups, Results "A B cell gene reporter panel").
# Names are the ones the paper prints in Fig. 2A and its text. Table S1 and S3 use the gene symbols in ALIAS.
GROUPS = {
    "BCR signalling": ["BLK", "BLNK", "BTK", "CD20", "CD27", "CD45", "CD79A", "CD79B", "CD81", "CSK", "IGH", "IGK", "LYN", "PLCG2", "SYK"],
    "PI3K-AKT-mTOR": ["AKT2", "BCL-XL", "BCL2", "MCL1", "PIK3AP1", "PTEN"],
    "NF-kB signalling": ["MAP3K7", "TNIP1"],
    "Nuclear proteins and transcription factors": ["BACH2", "BCL11A", "BCL6", "BTG2", "E2F2", "HDAC1", "ID3", "IKZF1", "IRF8", "MEF2B", "MSH2", "MSH6",
                                                   "MYC", "NCOA3", "OCAB", "OCT2", "PARP1", "PAX5", "RB1", "REL", "SPI1", "SPIB", "TCF3", "TP53"],
}
ALIAS = {"OCAB": ("OcaB", "POU2AF1"), "IGK": ("IgKappa", "IGK"), "IGH": ("IgH", "IGH"), "CD45": ("PTPRC", "PTPRC"), "CD20": ("MS4A1", "MS4A1"), "OCT2": ("POU2F2", "POU2F2")}

# genes the paper names as changed by rapamycin (Results "mTOR suppresses immune response genes", paragraphs 1 and 3)
SETS = [("Group 1 genes", ["MSH2", "MSH6", "E2F2"]), ("Immune response genes", ["IL4R", "CD86", "CD79A", "CD79B", "BACH2", "BCL6", "AICDA"])]


def rows(path: Path, sheet=0, skip=3):
    ws = openpyxl.load_workbook(path, read_only=True)
    ws = ws[sheet] if isinstance(sheet, str) else ws.worksheets[sheet]
    return [r for r in list(ws.iter_rows(values_only=True))[skip:] if r[0] is not None]


def deg(path: Path) -> dict:
    d = {r[0]: r for r in rows(path)}
    return d   # name -> (name, baseMean, log2FoldChange, lfcSE, stat, pvalue, padj)


def must(name: str, got, want):
    assert got == want, f"{name}: table gives {got}, the paper's text says {want}"


def main() -> None:
    OUT.mkdir(exist_ok=True)
    f = lambda n: DATA / f"NIHMS2120106-supplement-{n}.xlsx"   # noqa: E731

    # ---- Table S1: the 47 reporter genes
    s1 = {}
    ws = openpyxl.load_workbook(f(5), read_only=True).worksheets[0]
    kind = term = None
    for r in list(ws.iter_rows(values_only=True))[2:]:
        if r[0]:
            kind = r[0].split()[0]       # mOrange or GFP
        if r[1]:
            term = r[1]
        if r[2] and r[2] != "Gene":
            s1[r[2]] = (kind, term)
    must("Table S1 genes", len(s1), 47)

    # ---- Table S3: regulators and links per screen
    wb = openpyxl.load_workbook(f(7), read_only=True)
    A = [r for r in list(wb["Table S3A"].iter_rows(values_only=True))[3:] if r[0] is not None]
    R = [r for r in list(wb["Table S3B"].iter_rows(values_only=True))[3:] if r[0] is not None]
    sa, sr = {r[1] for r in A}, {r[1] for r in R}
    must("activating links", len(A), 6487)           # Methods, "Gene regulatory network"
    must("repressing links", len(R), 11151)
    must("links", len(A) + len(R), 17638)            # Results, "A comprehensive B cell regulatory network"
    must("regulators", len(sa | sr), 4440)
    must("activators only", len(sa - sr), 1685)
    must("both", len(sa & sr), 1002)
    must("repressors only", len(sr - sa), 1753)
    ca, cr = Counter(r[0] for r in A), Counter(r[0] for r in R)
    targets = sorted(set(ca) | set(cr))
    must("screens", len(targets), 47)
    must("MSH6 activators", ca["MSH6"], 376)         # text: "MSH6 was linked to 376 activators and 136 repressors"
    must("MSH6 repressors", cr["MSH6"], 136)
    must("BCL11A repressors", cr["BCL11A"], 298)
    must("BCL11A activators", ca["BCL11A"], 29)
    mean_a, mean_r = round(len(A) / 47), round(len(R) / 47)
    must("mean activators per screen", mean_a, 138)  # text: "approximately twice as many repressors (n = 237) as activators (n = 138)"
    must("mean repressors per screen", mean_r, 237)

    # ---- the 47 genes of the figure against Tables S1 and S3
    label_of = {}
    for g in [x for v in GROUPS.values() for x in v]:
        s1n, s3n = ALIAS.get(g, (g, g))
        label_of[g] = (s1n, s3n)
    must("figure genes", len(label_of), 47)
    must("Table S1 names", sorted(v[0] for v in label_of.values()), sorted(s1))
    must("Table S3 names", sorted(v[1] for v in label_of.values()), sorted(targets))
    with (OUT / "reporters_47.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["label_in_figure", "group_in_figure", "name_in_table_S1", "name_in_table_S3", "reporter", "terminus", "activators_table_S3A", "repressors_table_S3B"])
        for grp, gs in GROUPS.items():
            for g in gs:
                s1n, s3n = label_of[g]
                w.writerow([g, grp, s1n, s3n, s1[s1n][0], s1[s1n][1], ca[s3n], cr[s3n]])

    # ---- DEG tables
    S5, S9, S10 = deg(f(9)), deg(f(2)), deg(f(3))
    must("rapamycin genes changed", len(S5), 2458)   # "We identified 2,458 DEGs"
    must("LY2090314 genes changed", len(S9), 427)    # "identified 427 DEGs"
    down9 = [k for k, v in S9.items() if v[2] < 0]
    must("LY2090314 down", len(down9), 356)          # "n = 356, or 83%"
    must("LY2090314 down, share", round(100 * len(down9) / len(S9)), 83)
    opp = [k for k in down9 if k in S5 and S5[k][2] > 0]
    must("down with LY2090314 and up with rapamycin", len(opp), 248)   # "70% of genes downregulated by GSK3 inhibition (n = 248)"
    must("share", round(100 * len(opp) / len(down9)), 70)
    up5 = sum(1 for v in S5.values() if v[2] > 0)
    shared = [k for k in S9 if k in S5]
    quad = Counter(("rapa up" if S5[k][2] > 0 else "rapa down") + ", " + ("GSK3i up" if S9[k][2] > 0 else "GSK3i down") for k in shared)
    x = [S5[k][2] for k in shared]
    y = [S9[k][2] for k in shared]
    n = len(x)
    mx, my = sum(x) / n, sum(y) / n
    r = sum((a - mx) * (b - my) for a, b in zip(x, y)) / math.sqrt(sum((a - mx) ** 2 for a in x) * sum((b - my) ** 2 for b in y))
    with (OUT / "shared_303.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["gene", "log2FC_rapamycin_S5", "log2FC_LY2090314_S9"])
        for k in shared:
            w.writerow([k, S5[k][2], S9[k][2]])

    # ---- the genes the paper names
    s6 = openpyxl.load_workbook(f(10), read_only=True)
    clus = {}
    for sh in ("Table S6A", "Table S6B"):
        clus[sh] = {r[1]: r[0] for r in list(s6[sh].iter_rows(values_only=True))[3:] if r[1]}
    with (OUT / "example_genes.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["set", "gene", "log2FC_rapamycin_S5", "padj_S5", "log2FC_LY2090314_S9", "padj_S9", "log2FC_both_S10", "padj_S10", "STRING_cluster_S6A", "STRING_cluster_S6B"])
        for st, gs in SETS:
            for g in gs:
                row = [st, g]
                for T in (S5, S9, S10):
                    row += [T[g][2], T[g][6]] if g in T else ["", ""]
                row += [clus["Table S6A"].get(g, ""), clus["Table S6B"].get(g, "")]
                w.writerow(row)
    for g in SETS[0][1]:
        assert g in S5 and S5[g][2] < 0 and g in clus["Table S6A"], g       # down with rapamycin, in the STRING list of down genes
    for g in SETS[1][1]:
        assert g in S5 and S5[g][2] > 0 and g in clus["Table S6B"], g       # up with rapamycin, in the STRING list of up genes
    # ribosome and translation genes: the second cluster of the STRING list of genes raised by rapamycin (Table S6B), and what the two other tables list for them
    tr = sorted(g for g, c in clus["Table S6B"].items() if c == 2)
    must("translation cluster size", len(tr), 99)
    assert sum(1 for g in tr if g.startswith(("RPL", "RPS", "EEF", "EIF"))) >= 70, "cluster 2 of S6B is not the ribosome and translation cluster"
    assert all(g in S5 and S5[g][2] > 0 for g in tr), "every translation gene is raised in S5"
    tr_up10 = [g for g in tr if g in S10 and S10[g][2] > 0]
    tr_dn10 = [g for g in tr if g in S10 and S10[g][2] < 0]
    must("translation genes listed up with both drugs", len(tr_up10), 74)
    must("translation genes listed down with both drugs", len(tr_dn10), 0)
    cut = min(min(abs(S5[k][2]), abs(S9[k][2])) for k in shared)
    assert abs(cut - math.log2(1.5)) < 0.01, "Methods: genes changed by more than 1.5-fold (the smallest in the tables is 1.495-fold)"
    rep_up = sorted(g for g in label_of if label_of[g][1] in S5 and S5[label_of[g][1]][2] > 0)
    rep_dn = sorted(g for g in label_of if label_of[g][1] in S5 and S5[label_of[g][1]][2] < 0)

    counts = {
        "reporter_lines": 47, "screens": 47, "regulators": 4440, "activators_only": 1685, "both_roles": 1002, "repressors_only": 1753,
        "links": 17638, "activating_links": 6487, "repressing_links": 11151, "mean_activators_per_screen": mean_a, "mean_repressors_per_screen": mean_r,
        "rapamycin_genes_changed": len(S5), "rapamycin_up": up5, "rapamycin_down": len(S5) - up5,
        "gsk3i_genes_changed": len(S9), "gsk3i_down": len(down9), "gsk3i_up": len(S9) - len(down9),
        "both_genes_changed_S10": len(S10),
        "translation_genes_S6B_cluster2": len(tr), "translation_up_with_both_S10": len(tr_up10), "smallest_abs_log2FC_of_the_303": round(cut, 3),
        "shared_genes": len(shared), "quadrants": dict(quad), "pearson_r_shared": round(r, 3),
        "opposite_direction": quad["rapa up, GSK3i down"] + quad["rapa down, GSK3i up"],
        "reporters_up_with_rapamycin_S5": len(rep_up), "reporters_down_with_rapamycin_S5": len(rep_dn),
        "source": {"reporters": "Table S1", "regulators and links": "Table S3A, S3B", "rapamycin": "Table S5", "LY2090314": "Table S9", "both": "Table S10", "clusters": "Table S6A, S6B"},
    }
    (OUT / "counts.json").write_text(json.dumps(counts, indent=1) + "\n", encoding="utf-8")
    print(json.dumps({k: v for k, v in counts.items() if k != "source"}))
    print("reporter genes up with rapamycin (S5):", len(rep_up), rep_up)
    print("reporter genes down with rapamycin (S5):", len(rep_dn), rep_dn)


if __name__ == "__main__":
    main()
