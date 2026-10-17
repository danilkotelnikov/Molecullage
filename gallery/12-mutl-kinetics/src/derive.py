"""Numbers of gallery piece 12 (redesign of Fig. 1D and E of Fukui et al., eLife 15:RP111443, doi 10.7554/eLife.111443, CC BY 4.0).

    py -3 -B -I gallery/12-mutl-kinetics/src/derive.py          # run fetch_data.py first

Reads data/elife-111443-fig1-data2-v1.xlsx ("Figure 1 - source data 2": mean and SD of kapp against [ATP], three independent experiments)
and Table 1 of data/PMC13614785.xml, checks them against each other, and writes
  tables/table1_mutl.csv       the 12 aqMutL rows of Table 1 exactly as printed (strings), plus the rank used on the page
  tables/kapp_points.csv       every cell of the source file: variant, [ATP] (mM), mean kapp, SD (blank where the file is blank)
  tables/refit_vs_table1.csv   my least-squares refit of the five variants Table 1 gives numbers for, set against the printed values
  tables/not_determined.csv    for the seven N.D. variants: the highest mean kapp in the file and the [ATP] range measured
  tables/derived.json          everything make_layer.py draws (no third-party file is read there)
Nothing in the figure is smoothed: the curves are Michaelis-Menten curves with the parameters printed in Table 1 (as the paper's own curves are),
the dashed curves are my unweighted least-squares refit of the plotted means (scipy curve_fit, kapp = kcat [ATP] / (Km + [ATP])).
"""
from __future__ import annotations

import csv
import json
import math
import sys
from pathlib import Path

import numpy as np
import openpyxl
from scipy.optimize import curve_fit

HERE = Path(__file__).resolve().parent
GAL = HERE.parent
sys.path.insert(0, str(GAL))
from fetch_data import read_table1  # noqa: E402

XLSX = GAL / "data" / "elife-111443-fig1-data2-v1.xlsx"
XML = GAL / "data" / "PMC13614785.xml"
TAB = GAL / "tables"
TAB.mkdir(exist_ok=True)

# sheet column header -> name in the paper (Table 1 notation)
SHEET_NAME = {"Wildtype": "WT", "E29Q_E32Q": "E29Q/E32Q"}
MUTL_ROWS = ["WT", "E29A", "E29V", "E29K", "E29Q", "E29G", "E32A", "E32V", "E32K", "E32Q", "E32G", "E29Q/E32Q"]   # Table 1 order, aqMutL block
CLASS = lambda n: "wt" if n == "WT" else ("both" if "/" in n else ("g29" if n.startswith("E29") else "g32"))


def mm(s, kcat, km):
    return kcat * s / (km + s)


def read_table():
    rows = read_table1(XML)
    out = {}
    for r in rows[:12]:
        name = "WT" if r[0] == "aqMutL WT" else r[0]
        out[name] = r
    assert list(out) == MUTL_ROWS, list(out)
    assert rows[12][0] == "aqGyrB NTD WT", rows[12]       # the MutL block is rows 1-12; GyrB, PMS2 and MLH1 follow and are not drawn
    res = {}
    for n, (_, kc, km, kk) in out.items():
        if kc.startswith("N.D."):
            assert km.startswith("N.D.") and kk.startswith("N.D."), n
            res[n] = dict(nd=True)
        else:
            (a, ea), (b, eb) = kc.split("±"), km.split("±")
            res[n] = dict(nd=False, kcat_s=a, kcat_e=ea, km_s=b, km_e=eb, kcatkm_s=kk, kcat=float(a), kcat_err=float(ea), km=float(b), km_err=float(eb), kcatkm=float(kk))
            # the printed kcat/Km is kcat [min-1] / Km [mM] * 1000 / 60 (M-1 s-1), within rounding of the printed digits
            calc = float(a) / float(b) * 1000 / 60
            assert abs(calc - float(kk)) / float(kk) < 0.03, (n, calc, kk)
    return res


def read_points():
    ws = openpyxl.load_workbook(XLSX, data_only=True).active
    assert ws["A2"].value == "[ATP] (mM)" and ws["J2"].value == "[ATP] (mM)"
    f2 = openpyxl.load_workbook(XLSX, data_only=False).active
    col = lambda c, r0, r1: [ws[f"{c}{r}"].value for r in range(r0, r1 + 1)]
    out, notes = {}, []
    for hdr_row, xcol, cols, rows_mean, rows_sd, label in ((2, "A", "BCDEFGH", (3, 14), (16, 27), "left block"), (2, "J", "KLMNO", (3, 10), (12, 19), "right block")):
        xm, xs = col(xcol, *rows_mean), col(xcol, *rows_sd)
        assert xm == xs, (label, xm, xs)                       # the SD rows repeat the [ATP] ladder of the mean rows
        assert ws[f"{xcol}{rows_mean[1] + 1}"].value == "SD" and rows_sd[0] == rows_mean[1] + 2      # the label SD sits between the two blocks of rows
        for c in cols:
            raw = ws[f"{c}{hdr_row}"].value
            name = SHEET_NAME.get(raw, raw)
            mean, sd = col(c, *rows_mean), col(c, *rows_sd)
            out[name] = [dict(x=float(x), mean=(None if m is None else float(m)), sd=(None if s is None else float(s))) for x, m, s in zip(xm, mean, sd)]
            for r in range(rows_mean[0], rows_mean[1] + 1):
                v = f2[f"{c}{r}"].value
                if isinstance(v, str) and v.startswith("="):
                    notes.append(f"{name}: cell {c}{r} holds the formula {v}; its stored value {ws[f'{c}{r}'].value:.6g} is used")
    assert set(out) == set(MUTL_ROWS), set(out) ^ set(MUTL_ROWS)
    return out, notes


def fit(pts):
    x = np.array([p["x"] for p in pts if p["mean"] is not None])
    y = np.array([p["mean"] for p in pts if p["mean"] is not None])
    p, cov = curve_fit(mm, x, y, p0=[1.5 * y.max(), 1.0], maxfev=20000)
    se = np.sqrt(np.diag(cov))
    return dict(kcat=float(p[0]), kcat_se=float(se[0]), km=float(p[1]), km_se=float(se[1]), r=float(cov[0, 1] / se[0] / se[1]), n=int(len(x)))


def fmt(v, nd=2):
    return f"{v:.{nd}g}"


def main():
    T = read_table()
    P, notes = read_points()
    det = [n for n in MUTL_ROWS if not T[n]["nd"]]
    nd = [n for n in MUTL_ROWS if T[n]["nd"]]
    assert det == ["WT", "E29Q", "E32A", "E32V", "E32Q"] and nd == ["E29A", "E29V", "E29K", "E29G", "E32K", "E32G", "E29Q/E32Q"], (det, nd)
    rank = sorted(det, key=lambda n: -T[n]["kcatkm"]) + nd          # ranked by the printed kcat/Km, high to low; N.D. variants last, in Table 1 order
    assert rank[:5] == ["WT", "E32A", "E29Q", "E32Q", "E32V"], rank

    # ---- the units of kapp: the file states none. Table 1 gives kcat in min-1, the original panels D and E are labelled min-1; the file's wild-type plateau
    # (0.58 at 1 mM) is below the Table 1 kcat 0.65 min-1 and the refit agrees, so the file is read as min-1.
    wt_top = max(p["mean"] for p in P["WT"])
    assert 0.5 < wt_top < T["WT"]["kcat"] * 1.05, wt_top

    # ---- refit of the five variants Table 1 has numbers for: unweighted least squares on the means, (0, 0) included
    R = {n: fit(P[n]) for n in det}
    for n in det:
        R[n]["kcatkm"] = R[n]["kcat"] / R[n]["km"] * 1000 / 60
    # the pair E32A / E32Q: the refit of each file column matches the Table 1 row of the OTHER name (Km 2.1 against 3.5 mM and the reverse)
    sw = {}
    for a, b in (("E32A", "E32Q"),):
        km_ta, km_tb = T[a]["km"], T[b]["km"]
        own = abs(R[a]["km"] / km_ta - 1) + abs(R[b]["km"] / km_tb - 1)
        crossed = abs(R[a]["km"] / km_tb - 1) + abs(R[b]["km"] / km_ta - 1)
        sw[f"{a}/{b}"] = dict(own_relative_difference=own, crossed_relative_difference=crossed)
        assert crossed < 0.06 and own > 0.5, (own, crossed)          # crossed within 6 % (Km 3.48 vs 3.5 and 2.12 vs 2.1); own labels more than 50 % off
        # the whole printed row follows the crossing: kcat, the Km error, and the kcat error (printed error / refit standard error)
        sw[f"{a}/{b}"]["file_column_E32A_vs_table_row_E32Q"] = dict(kcat=(R[a]["kcat"], T[b]["kcat"]), km=(R[a]["km"], T[b]["km"]), km_err=(R[a]["km_se"], T[b]["km_err"]),
                                                                   kcat_err_over_se=T[b]["kcat_err"] / R[a]["kcat_se"])
        sw[f"{a}/{b}"]["file_column_E32Q_vs_table_row_E32A"] = dict(kcat=(R[b]["kcat"], T[a]["kcat"]), km=(R[b]["km"], T[a]["km"]), km_err=(R[b]["km_se"], T[a]["km_err"]),
                                                                   kcat_err_over_se=T[a]["kcat_err"] / R[b]["kcat_se"])
    for n in ("WT", "E29Q", "E32V"):
        assert abs(R[n]["km"] / T[n]["km"] - 1) < 0.05 and abs(R[n]["kcat"] / T[n]["kcat"] - 1) < 0.08, (n, R[n], T[n])
    with open(TAB / "refit_vs_table1.csv", "w", newline="", encoding="utf-8") as fh:
        fh.write("# Refit of the source file (Figure 1 - source data 2): unweighted least squares of kapp = kcat [ATP] / (Km + [ATP]) to the means (scipy curve_fit, (0, 0) included),\n"
                 "# against Table 1 of Fukui et al., eLife 15:RP111443 as printed. se = standard error of the fit. kcat in min-1, Km in mM, kcat/Km in M-1 s-1 (kcat/Km x 1000/60).\n"
                 "# E32A and E32Q: the file's two columns fit the OTHER name's Table 1 row; the file and Table 1 disagree and the figure draws the file as labelled.\n")
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["variant", "n_points", "refit_kcat", "refit_kcat_se", "refit_Km", "refit_Km_se", "refit_kcat_over_Km", "corr_kcat_Km", "table_kcat", "table_kcat_err", "table_Km", "table_Km_err", "table_kcat_over_Km"])
        for n in det:
            r, t = R[n], T[n]
            w.writerow([n, r["n"], f"{r['kcat']:.4f}", f"{r['kcat_se']:.4f}", f"{r['km']:.4f}", f"{r['km_se']:.4f}", f"{r['kcatkm']:.3f}", f"{r['r']:.3f}", t["kcat_s"], t["kcat_e"], t["km_s"], t["km_e"], t["kcatkm_s"]])
    kcat_err_ratio = {n: (T[n]["kcat_err"] / R[n]["kcat_se"]) for n in det}

    # ---- not-determined variants: the highest mean kapp in the file, and the [ATP] range measured
    ndrows = []
    nd_info = {}
    for n in nd:
        pts = [p for p in P[n] if p["mean"] is not None]
        top = max(pts, key=lambda p: p["mean"])
        miss = [p["x"] for p in P[n] if p["mean"] is None]
        nd_info[n] = dict(max_mean=top["mean"], max_at=top["x"], x_max=max(p["x"] for p in pts), n_points=len(pts), missing_x=miss)
        ndrows.append([n, f"{top['mean']:.6f}", top["x"], max(p["x"] for p in pts), len(pts), ";".join(f"{m:g}" for m in miss)])
    with open(TAB / "not_determined.csv", "w", newline="", encoding="utf-8") as fh:
        fh.write("# The seven variants Table 1 marks N.D. (activity not detected). Highest mean kapp (min-1) of the 8 or 12 [ATP] points of the source file (Figure 1 - source data 2),\n"
                 "# the [ATP] where it occurs, the highest [ATP] with a value, the number of points with a value, and [ATP] values (mM) where the file cell is blank.\n")
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["variant", "highest_mean_kapp_per_min", "at_ATP_mM", "highest_ATP_mM", "points_with_value", "blank_cells_at_ATP_mM"])
        w.writerows(ndrows)

    # ---- every cell of the file
    with open(TAB / "kapp_points.csv", "w", newline="", encoding="utf-8") as fh:
        fh.write("# Figure 1 - source data 2 of Fukui et al., eLife 15:RP111443 (CC BY 4.0): mean and SD of kapp (the file gives no unit; read as min-1, see how.md) of three independent experiments.\n"
                 "# Variant names are the paper's (Table 1); the file's Wildtype is WT and E29Q_E32Q is E29Q/E32Q. Blank = the cell is blank in the file.\n")
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["variant", "ATP_mM", "mean_kapp", "sd_kapp"])
        for n in MUTL_ROWS:
            for p in P[n]:
                w.writerow([n, f"{p['x']:g}", "" if p["mean"] is None else f"{p['mean']:.8g}", "" if p["sd"] is None else f"{p['sd']:.8g}"])

    # ---- Table 1 as printed, with the rank
    with open(TAB / "table1_mutl.csv", "w", newline="", encoding="utf-8") as fh:
        fh.write("# Table 1 of Fukui et al., eLife 15:RP111443 (CC BY 4.0), the 12 aqMutL rows, as printed in the Europe PMC full text (PMC13614785). N.D. = the activity was not detected (table footnote).\n"
                 "# kcat in min-1, Km in mM, kcat/Km in M-1 s-1; the table prints no error for kcat/Km. rank: by kcat/Km high to low, N.D. rows last in table order.\n")
        w = csv.writer(fh, lineterminator="\n")
        w.writerow(["rank", "variant", "kcat", "kcat_err", "Km", "Km_err", "kcat_over_Km"])
        for i, n in enumerate(rank, 1):
            t = T[n]
            w.writerow([i, n] + (["N.D.", "", "N.D.", "", "N.D."] if t["nd"] else [t["kcat_s"], t["kcat_e"], t["km_s"], t["km_e"], t["kcatkm_s"]]))

    # ---- what the page draws
    wt = T["WT"]["kcatkm"]
    ratios = {n: wt / T[n]["kcatkm"] for n in det if n != "WT"}      # computed from the printed (rounded) kcat/Km values
    assert [round(ratios[n]) for n in ("E32A", "E29Q", "E32Q", "E32V")] == [97, 100, 160, 171], ratios
    var = {}
    for n in MUTL_ROWS:
        t = T[n]
        v = dict(name=n, cls=CLASS(n), nd=t["nd"], pts=[p for p in P[n]])
        if t["nd"]:
            v.update(nd_info[n])
        else:
            v.update({k: t[k] for k in ("kcat", "kcat_err", "km", "km_err", "kcatkm", "kcat_s", "kcat_e", "km_s", "km_e", "kcatkm_s")})
            v["refit"] = R[n]
            v["xmax"] = max(p["x"] for p in P[n])
            v["ratio"] = ratios.get(n)
        var[n] = v
    D = dict(rank=rank, det=det, nd=nd, variants=var, wt_top=wt_top, wt_top_at=max(P["WT"], key=lambda p: p["mean"])["x"], ratios={k: round(v, 1) for k, v in ratios.items()},
             ratio_min=round(min(ratios.values()), 1), ratio_max=round(max(ratios.values()), 1), swap=sw, notes=notes,
             kcat_err_over_fit_se={k: round(v, 2) for k, v in kcat_err_ratio.items()})
    (TAB / "derived.json").write_text(json.dumps(D, indent=1), encoding="utf-8")

    print("ranking:", rank)
    print("ratios wild type / variant (computed from printed kcat/Km):", {k: round(v, 1) for k, v in ratios.items()})
    for n in det:
        r, t = R[n], T[n]
        print(f"{n:5s} refit kcat {r['kcat']:.4f} Km {r['km']:.3f} kcat/Km {r['kcatkm']:.3f} | Table kcat {t['kcat_s']}±{t['kcat_e']} Km {t['km_s']}±{t['km_e']} kcat/Km {t['kcatkm_s']} | Km se {r['km_se']:.3f} corr {r['r']:.2f}")
    print("E32A/E32Q crossing evidence:", json.dumps(sw))
    print("Table kcat error / refit kcat se:", {k: round(v, 2) for k, v in kcat_err_ratio.items()})
    for n in nd:
        print(f"{n:10s} highest mean kapp {nd_info[n]['max_mean']:.5f} at {nd_info[n]['max_at']:g} mM; measured to {nd_info[n]['x_max']:g} mM; blank at {nd_info[n]['missing_x']}")
    for s in notes:
        print("note:", s)
    print("WT highest mean kapp", round(wt_top, 4), "at", D["wt_top_at"], "mM")


if __name__ == "__main__":
    main()
